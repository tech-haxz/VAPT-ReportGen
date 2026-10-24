import tempfile, zipfile, io
from fastapi import HTTPException
from starlette.datastructures import UploadFile
from PIL import Image
from docx import Document
from app import main

def setup_workspace(td):
    main.DB=main.Path(td)/'test.sqlite3'; main.DATA=main.Path(td); main.UPLOADS=main.DATA/'evidence'; main.REPORTS=main.DATA/'reports'; main.UPLOADS.mkdir(); main.REPORTS.mkdir(); main.init()

def test_assessment_finding_report_workflow():
    with tempfile.TemporaryDirectory(dir=main.ROOT/'backend'/'tests') as td:
        setup_workspace(td)
        a=main.create_assessment(main.Assessment(title='Fictional Demo Assessment',target='demo.example.test')); aid=a['id']
        f=main.create_finding(aid,main.Finding(title='Demo missing security header',location='https://demo.example.test/search?q=test',evidence='Fictional HTTP response lacked the example header.',severity='High',cve='CVE-TEST-0001',cwe='CWE-TEST',references_text='Bugtraq ID TEST-1',method='Fictional scanner 1.0',remediation='Configure the header.')); fid=f['id']
        assert main.get_finding(fid)['severity']=='High'
        image=io.BytesIO(); Image.new('RGB',(2,2),(35,180,160)).save(image,format='PNG'); png=image.getvalue()
        up=main.upload_evidence(fid,UploadFile(filename='fictional.png',file=io.BytesIO(png)),caption='Fictional response placeholder')
        assert up['sha256'] and main.get_finding(fid)['evidence_files']
        assert main.update_finding(fid,main.Finding(**{**f,'reviewed':True}))['reviewed']==1
        rep=main.export_report(aid,{'kind':'consolidated'}); path=main.REPORTS/(rep['id']+'.docx')
        assert path.read_bytes()[:2]==b'PK'
        download=main.download_report(rep['id'])
        assert main.Path(download.path)==path
        assert download.media_type=='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        assert rep['filename'] in download.headers['content-disposition']
        with zipfile.ZipFile(path) as z: assert any(x.startswith('word/media/') for x in z.namelist())
        report=Document(path); paragraphs='\n'.join(p.text for p in report.paragraphs)
        assert [cell.text for cell in report.tables[0].rows[0].cells]==[
            'S. No.','Vulnerable location, parameter, path, link','Name of the vulnerability',
            'References (CVE-ID, CWE-ID, Bugtraq-ID, etc.)']
        row=[cell.text for cell in report.tables[0].rows[1].cells]
        assert row[0]=='1.' and 'https://demo.example.test/search?q=test' in row[1]
        assert 'CVE-TEST-0001' in row[3] and 'CWE-TEST' in row[3] and 'Bugtraq ID TEST-1' in row[3]
        assert 'Table 1. Vulnerability Summary' in paragraphs
        assert 'Other Information' in paragraphs and 'Fictional scanner 1.0' in paragraphs and 'Part 2: Detailed Report' in paragraphs
        assert 'Location of Vulnerability / Error' in paragraphs
        assert 'Proof of concept and steps of verification' in paragraphs
        assert 'Workaround / Solutions' in paragraphs and 'Configure the header.' in paragraphs
        assert 'Executive summary' not in paragraphs and 'Severity distribution' not in paragraphs
        assert main.delete_report(rep['id'])=={'deleted':True,'already_deleted':False}
        assert not path.exists()
        assert all(r['id']!=rep['id'] for r in main.reports())
        assert main.delete_report(rep['id'])=={'deleted':False,'already_deleted':True}
        with main.connect() as db:
            stored=db.execute('SELECT stored FROM evidence WHERE finding_id=?',(fid,)).fetchone()['stored']
        evidence_path=main.UPLOADS/stored
        assert evidence_path.exists() and not path.exists()
        main.delete_assessment(aid)
        assert not evidence_path.exists() and not path.exists()
        with main.connect() as db:
            assert db.execute('SELECT COUNT(*) FROM findings WHERE assessment_id=?',(aid,)).fetchone()[0]==0
            assert db.execute('SELECT COUNT(*) FROM reports WHERE assessment_id=?',(aid,)).fetchone()[0]==0

def test_manual_ai_and_bad_upload():
    with tempfile.TemporaryDirectory(dir=main.ROOT/'backend'/'tests') as td:
        setup_workspace(td)
        a=main.create_assessment(main.Assessment(title='Demo')); f=main.create_finding(a['id'],main.Finding(title='Test',location='/',description='Observed',evidence='Response'))
        assert main.analyze(f['id'])['available'] is False
        try: main.upload_evidence(f['id'],UploadFile(filename='bad.png',file=io.BytesIO(b'not-image')))
        except HTTPException as e: assert e.status_code==400
        else: assert False,'Invalid image should be rejected'



